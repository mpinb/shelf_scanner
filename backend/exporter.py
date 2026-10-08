import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def generate_library_xlsx(books_data) -> bytes:
    """
    Generates a beautifully styled Excel (.xlsx) workbook for a list of book records.
    Ensures safe character encoding, auto-fitted columns, frozen headers, and text-safe ISBNs.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "My Library Catalog"

    headers = [
        "Shelf", "Index", "Detected Title", "OpenLibrary Canonical Title",
        "Authors", "OpenLibrary Authors", "Publisher", "OpenLibrary Publisher",
        "Pub Year", "Primary ISBN", "All Potential ISBNs", "Subjects / Tags",
        "Edition / Misc", "Spine OCR Text", "User Notes", "Is Ignored", "Enrichment Status"
    ]
    ws.append(headers)

    # 1. Header Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
    ws.row_dimensions[1].height = 26

    # 2. Data Rows Styling
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    regular_font = Font(name="Calibri", size=10)
    isbn_font = Font(name="Consolas", size=10, color="047857")

    def clean_val(v):
        if v is None:
            return ""
        if isinstance(v, (int, float)):
            return v
        s = ILLEGAL_CHARACTERS_RE.sub("", str(v).strip())
        # Prevent Formula / CSV Injection in Excel:
        # If string starts with =, +, -, @, or tab/carriage return, prepend with single quote
        if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
            s = f"'{s}"
        return s

    for row_idx, b in enumerate(books_data, start=2):
        item = dict(b)
        isbns_list = []
        isbns_val = item.get("isbns")
        if isinstance(isbns_val, list):
            isbns_list = isbns_val
        elif isinstance(isbns_val, str) and isbns_val.strip():
            try:
                isbns_list = json.loads(isbns_val)
            except Exception:
                isbns_list = [isbns_val]
        isbns_str = "; ".join(isbns_list)

        shelf_name = item.get("shelf_name") or item.get("shelf_id") or ""

        row_values = [
            clean_val(shelf_name),
            clean_val(item.get("book_index", "")),
            clean_val(item.get("title", "") or ""),
            clean_val(item.get("ol_title", "") or ""),
            clean_val(item.get("authors", "") or ""),
            clean_val(item.get("ol_authors", "") or ""),
            clean_val(item.get("publication", "") or ""),
            clean_val(item.get("ol_publisher", "") or ""),
            clean_val(item.get("pub_year", "") or ""),
            clean_val(item.get("isbn_primary", "") or ""),
            clean_val(isbns_str),
            clean_val(item.get("subjects", "") or ""),
            clean_val(item.get("misc", "") or ""),
            clean_val(item.get("raw_text", "") or ""),
            clean_val(item.get("user_notes", "") or ""),
            "Yes" if item.get("is_ignored") else "No",
            clean_val(item.get("enrichment_status", "") or "pending")
        ]
        ws.append(row_values)

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # ISBN column: text format to avoid scientific notation
            if col_idx == 10:
                cell.number_format = '@'
                cell.font = isbn_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (1, 2, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
